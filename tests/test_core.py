from __future__ import annotations

import tempfile
import time
import unittest
from datetime import date
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor

from jarvis_bambu.discovery import discover_pieces
from jarvis_bambu.cli_controller import BambuStudioCliController
from jarvis_bambu.excel_report import ExcelReport
from jarvis_bambu.models import PieceMetrics
from jarvis_bambu.global_report import GlobalReportBuilder
from jarvis_bambu.optimizer_models import OptimizerOptions


class CoreWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.template = Path(__file__).resolve().parents[1] / "Plantilla_Analisis_Piezas.xlsx"
        self.defaults = {
            "machine": "P2S",
            "material": "PLA",
            "profile": "0.20 mm Standard",
            "filament_cost_per_gram": 0.0115,
            "machine_wear_per_hour": 0.20,
            "filament_sale_multiplier": 3.0,
        }

    def tearDown(self):
        self.temp.cleanup()

    def test_latest_and_today_discovery(self):
        first = self.root / "primera.stl"
        second = self.root / "segunda.stl"
        ignored = self.root / "nota.txt"
        first.write_bytes(b"one")
        time.sleep(0.02)
        ignored.write_text("x", encoding="utf-8")
        time.sleep(0.02)
        second.write_bytes(b"two")

        latest = discover_pieces(self.root, (".stl",), "latest")
        today = discover_pieces(self.root, (".stl",), "today", date.today())
        self.assertEqual(latest, [second])
        self.assertEqual({x.name for x in today}, {"primera.stl", "segunda.stl"})

    def test_filament_weight_from_length(self):
        controller = object.__new__(BambuStudioCliController)
        controller.defaults = {
            "filament_diameter_mm": 1.75,
            "filament_density_g_cm3": 1.26,
        }
        self.assertEqual(controller._grams_from_length(19.01), 57.61)

    def test_optimizer_effort_profiles_and_custom_minutes(self):
        low = OptimizerOptions("advanced", "low")
        high = OptimizerOptions("advanced", "high", preview=True)
        custom = OptimizerOptions("advanced", "3")
        self.assertEqual(low.time_limit_seconds, 60.0)
        self.assertIsNone(high.time_limit_seconds)
        self.assertTrue(high.preview)
        self.assertEqual(custom.time_limit_seconds, 180.0)
        self.assertEqual(custom.angle_step, 1)

    def test_excel_append_formula_image_and_duplicate(self):
        source = self.root / "pieza.stl"
        source.write_bytes(b"solid test")
        preview = self.root / "preview.png"
        Image.new("RGB", (320, 200), "#22aa66").save(preview)
        workbook_path = self.root / "report.xlsx"
        report = ExcelReport(workbook_path, self.template, self.defaults)
        row = report.append_success(
            PieceMetrics(
                source_path=source,
                print_minutes=125.0,
                weight_grams=100.0,
                preview_path=preview,
                machine="P2S",
                material="PLA",
                profile="0.20 mm Standard",
                sync_status="Sincronizado",
            )
        )
        self.assertEqual(row, 5)
        self.assertTrue(report.contains(source))

        workbook = load_workbook(workbook_path, data_only=False)
        sheet = workbook["Analisis"]
        self.assertEqual(sheet["B5"].value, "pieza")
        self.assertEqual(sheet["J5"].value, "=I5*'Configuracion'!$B$2")
        self.assertEqual(sheet["K5"].value, "=H5/60*'Configuracion'!$B$3")
        self.assertEqual(sheet["L5"].value, "=J5*'Configuracion'!$B$4+K5")
        self.assertEqual(len(sheet._images), 1)
        self.assertIsInstance(sheet._images[0].anchor, TwoCellAnchor)
        self.assertEqual(sheet._images[0].anchor.editAs, "twoCell")
        self.assertTrue(sheet.column_dimensions["N"].hidden)
        self.assertTrue(sheet["B5"].alignment.wrap_text)
        self.assertTrue(sheet["D5"].alignment.wrap_text)
        self.assertEqual(sheet.row_dimensions[5].height, 84)
        self.assertEqual(sheet.tables["AnalisisPiezas"].ref, "A4:O5")
        self.assertEqual(sheet.freeze_panes, "A5")

        second = self.root / "pieza_2.stl"
        second.write_bytes(b"solid second")
        second_row = report.append_success(
            PieceMetrics(
                source_path=second,
                print_minutes=30.0,
                weight_grams=25.0,
                preview_path=None,
                machine="P2S",
                material="PLA",
                profile="0.20 mm Standard",
            )
        )
        self.assertEqual(second_row, 6)
        workbook = load_workbook(workbook_path, data_only=False)
        self.assertEqual(workbook["Analisis"]["L6"].value, "=J6*'Configuracion'!$B$4+K6")
        self.assertEqual(workbook["Analisis"].tables["AnalisisPiezas"].ref, "A4:O6")

        output_folder = self.root / "daily"
        output_folder.mkdir()
        daily_path = output_folder / "Analisis_piezas_2026-08-26.xlsx"
        workbook_path.replace(daily_path)
        global_path = GlobalReportBuilder(
            output_folder,
            self.template,
            self.defaults,
            "Registro_global_piezas.xlsx",
        ).build()
        global_book = load_workbook(global_path, data_only=False)
        global_sheet = global_book["Analisis"]
        self.assertEqual(global_sheet["B5"].value, "pieza")
        self.assertEqual(global_sheet["B6"].value, "pieza_2")
        self.assertEqual(len(global_sheet._images), 1)
        self.assertEqual(global_sheet.tables["AnalisisPiezas"].ref, "A4:O6")


if __name__ == "__main__":
    unittest.main()
