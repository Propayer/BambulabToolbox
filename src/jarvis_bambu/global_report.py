from __future__ import annotations

import os
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

from .excel_report import ExcelReport, FIRST_DATA_ROW


class GlobalReportBuilder:
    """Rebuild the global workbook from every dated daily workbook."""

    def __init__(self, output_folder: Path, template_path: Path, defaults: dict, filename: str):
        self.output_folder = output_folder
        self.template_path = template_path
        self.defaults = defaults
        self.path = output_folder / filename

    def build(self) -> Path:
        self.output_folder.mkdir(parents=True, exist_ok=True)
        daily_files = sorted(self.output_folder.glob("Analisis_piezas_????-??-??.xlsx"))
        streams: list[BytesIO] = []

        handle, temporary_name = tempfile.mkstemp(
            prefix="registro_global_", suffix=".xlsx", dir=self.output_folder
        )
        os.close(handle)
        temporary_path = Path(temporary_name)
        temporary_path.unlink()

        try:
            report = ExcelReport(temporary_path, self.template_path, self.defaults)
            for daily_path in daily_files:
                self._append_workbook(report, daily_path, streams)
            report.save()
            temporary_path.replace(self.path)
        except Exception:
            temporary_path.unlink(missing_ok=True)
            raise
        return self.path

    def _append_workbook(
        self, report: ExcelReport, daily_path: Path, streams: list[BytesIO]
    ) -> None:
        workbook = load_workbook(daily_path, data_only=False)
        sheet = workbook["Analisis"]
        images_by_row: dict[int, list] = {}
        for image in sheet._images:
            if hasattr(image.anchor, "_from"):
                images_by_row.setdefault(image.anchor._from.row + 1, []).append(image)

        for source_row in range(FIRST_DATA_ROW, sheet.max_row + 1):
            if not any(sheet.cell(source_row, column).value for column in (1, 2, 4, 14)):
                continue
            target_row = report.first_empty_row()
            report._apply_row_style(target_row)
            for column in range(1, 16):
                if column in (10, 11, 12):
                    continue
                report.sheet.cell(target_row, column).value = sheet.cell(source_row, column).value

            report.sheet.cell(target_row, 10).value = (
                f"=I{target_row}*'Configuracion'!$B$2"
            )
            report.sheet.cell(target_row, 11).value = (
                f"=H{target_row}/60*'Configuracion'!$B$3"
            )
            report.sheet.cell(target_row, 12).value = (
                f"=J{target_row}*'Configuracion'!$B$4+K{target_row}"
            )

            for source_image in images_by_row.get(source_row, []):
                stream = BytesIO(source_image._data())
                streams.append(stream)
                report._add_cell_image(ExcelImage(stream), target_row)
            report._resize_table(target_row)

