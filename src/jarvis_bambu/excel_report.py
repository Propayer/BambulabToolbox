from __future__ import annotations

import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.worksheet.table import Table, TableStyleInfo

from .discovery import file_identity
from .models import PieceMetrics


HEADER_ROW = 4
FIRST_DATA_ROW = 5
MAX_DATA_ROW = 1004
TABLE_NAME = "AnalisisPiezas"
PRICE_COLUMNS = {
    "date": (1, "Fecha"), "name": (2, "Pieza"), "preview": (3, "Vista previa"),
    "file": (4, "Archivo"), "machine": (5, "Máquina"), "material": (6, "Material"),
    "profile": (7, "Perfil"), "time": (8, "Tiempo (min)"), "weight": (9, "Peso (g)"),
    "filament_cost": (10, "Coste filamento"), "wear": (11, "Desgaste"),
    "base_price": (12, "Venta aproximada"), "notes": (13, "Observaciones"),
    "sync": (15, "Sincronización"), "quantity": (16, "Cantidad"),
    "item_multiplier": (17, "Multiplicador pieza"),
    "global_multiplier": (18, "Multiplicador global"), "final_price": (19, "Precio final"),
}
MANDATORY_PRICE_COLUMNS = frozenset({"name"})
DEFAULT_PRICE_COLUMNS = (
    "date", "name", "file", "machine", "material", "profile", "time", "weight",
    "filament_cost", "wear", "base_price", "quantity", "item_multiplier",
    "global_multiplier", "final_price", "notes",
)


class ExcelReport:
    def __init__(self, workbook_path: Path, template_path: Path, defaults: dict):
        self.path = workbook_path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            shutil.copy2(template_path, self.path)
        self.workbook = load_workbook(self.path)
        self.sheet = self.workbook["Analisis"]
        self.config_sheet = self.workbook["Configuracion"]
        self.sheet.column_dimensions["N"].hidden = True
        self._write_defaults(defaults)
        self._normalize_existing_rows()
        self._ensure_table()
        self.sheet.freeze_panes = "A5"

    def _write_defaults(self, defaults: dict) -> None:
        self.config_sheet["B2"] = float(defaults["filament_cost_per_gram"])
        self.config_sheet["B3"] = float(defaults["machine_wear_per_hour"])
        self.config_sheet["B4"] = float(defaults["filament_sale_multiplier"])
        self.config_sheet["B5"] = str(defaults["machine"])
        self.config_sheet["B6"] = str(defaults["material"])
        self.config_sheet["B7"] = str(defaults["profile"])

    def known_identities(self) -> set[str]:
        return {
            str(self.sheet.cell(row=row, column=14).value)
            for row in range(FIRST_DATA_ROW, MAX_DATA_ROW + 1)
            if self.sheet.cell(row=row, column=14).value
        }

    def contains(self, path: Path) -> bool:
        return file_identity(path) in self.known_identities()

    def first_empty_row(self) -> int:
        for row in range(FIRST_DATA_ROW, MAX_DATA_ROW + 1):
            if not any(
                self.sheet.cell(row=row, column=column).value
                for column in (1, 2, 4, 14)
            ):
                return row
        raise RuntimeError("El Excel ha alcanzado el límite de 1.000 piezas.")

    def append_success(self, metrics: PieceMetrics) -> int:
        row = self.first_empty_row()
        self._apply_row_style(row)
        source = metrics.source_path
        values = {
            1: datetime.now(),
            2: source.stem,
            4: str(source),
            5: metrics.machine,
            6: metrics.material,
            7: metrics.profile,
            8: float(metrics.print_minutes),
            9: float(metrics.weight_grams),
            13: metrics.notes,
            14: file_identity(source),
            15: metrics.sync_status,
        }
        for column, value in values.items():
            self.sheet.cell(row=row, column=column).value = value

        self.sheet.cell(row=row, column=10).value = (
            f"=I{row}*'Configuracion'!$B$2"
        )
        self.sheet.cell(row=row, column=11).value = (
            f"=H{row}/60*'Configuracion'!$B$3"
        )
        self.sheet.cell(row=row, column=12).value = (
            f"=J{row}*'Configuracion'!$B$4+K{row}"
        )

        if metrics.preview_path and metrics.preview_path.exists():
            image = ExcelImage(str(metrics.preview_path))
            self._add_cell_image(image, row)

        self._resize_table(row)

        self.save()
        return row

    def append_error(self, source_path: Path, message: str, screenshot: Path | None) -> int:
        row = self.first_empty_row()
        self._apply_row_style(row)
        self.sheet.cell(row=row, column=1).value = datetime.now()
        self.sheet.cell(row=row, column=2).value = source_path.stem
        self.sheet.cell(row=row, column=4).value = str(source_path)
        self.sheet.cell(row=row, column=13).value = f"ERROR: {message}"
        self.sheet.cell(row=row, column=14).value = file_identity(source_path)
        self.sheet.cell(row=row, column=15).value = "Error"
        if screenshot and screenshot.exists():
            image = ExcelImage(str(screenshot))
            self._add_cell_image(image, row)
        self._resize_table(row)
        self.save()
        return row

    def append_priced_success(
        self, metrics: PieceMetrics, quantity: int, item_multiplier: float,
        global_multiplier: float,
    ) -> int:
        row = self.append_success(metrics)
        for column in range(16, 20):
            source = self.sheet.cell(row=row, column=12)
            target = self.sheet.cell(row=row, column=column)
            target._style = copy(source._style)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
        self.sheet.cell(row=HEADER_ROW, column=16).value = "Cantidad"
        self.sheet.cell(row=HEADER_ROW, column=17).value = "Multiplicador pieza"
        self.sheet.cell(row=HEADER_ROW, column=18).value = "Multiplicador global"
        self.sheet.cell(row=HEADER_ROW, column=19).value = "Precio final"
        self.sheet.cell(row=row, column=16).value = int(quantity)
        self.sheet.cell(row=row, column=17).value = float(item_multiplier)
        self.sheet.cell(row=row, column=18).value = float(global_multiplier)
        self.sheet.cell(row=row, column=19).value = f"=L{row}*P{row}*Q{row}*R{row}"
        self.sheet.cell(row=row, column=19).number_format = '#,##0.00 [$€-es-ES]'
        self.sheet.tables[TABLE_NAME].ref = f"A{HEADER_ROW}:S{max(row, FIRST_DATA_ROW)}"
        self.save()
        return row

    def configure_price_columns(self, selected: list[str]) -> None:
        visible = set(selected) | set(MANDATORY_PRICE_COLUMNS)
        for key, (column, _label) in PRICE_COLUMNS.items():
            self.sheet.column_dimensions[self.sheet.cell(HEADER_ROW, column).column_letter].hidden = key not in visible
        self.sheet.column_dimensions["N"].hidden = True
        self.save()

    def save(self) -> None:
        self.workbook.save(self.path)

    def _apply_row_style(self, row: int) -> None:
        if row != FIRST_DATA_ROW:
            for column in range(1, 16):
                source = self.sheet.cell(row=FIRST_DATA_ROW, column=column)
                target = self.sheet.cell(row=row, column=column)
                target._style = copy(source._style)
                if source.has_style:
                    target.font = copy(source.font)
                    target.fill = copy(source.fill)
                    target.border = copy(source.border)
                    target.alignment = copy(source.alignment)
                    target.number_format = source.number_format
                    target.protection = copy(source.protection)
        self._format_row(row)

    def _normalize_existing_rows(self) -> None:
        for row in range(FIRST_DATA_ROW, MAX_DATA_ROW + 1):
            if any(self.sheet.cell(row=row, column=column).value for column in (1, 2, 4, 14)):
                self._format_row(row)

        for image in self.sheet._images:
            anchor = image.anchor
            if hasattr(anchor, "_from"):
                row = anchor._from.row + 1
                self._set_cell_anchor(image, row)

    def _format_row(self, row: int) -> None:
        self.sheet.row_dimensions[row].height = 84
        for column in range(1, 16):
            cell = self.sheet.cell(row=row, column=column)
            alignment = copy(cell.alignment)
            alignment.vertical = "center"
            alignment.wrap_text = column in (2, 4, 13)
            alignment.shrink_to_fit = column in (5, 6, 7, 8, 9, 10, 11, 12, 15)
            cell.alignment = alignment

    def _last_data_row(self) -> int:
        rows = [
            row
            for row in range(FIRST_DATA_ROW, MAX_DATA_ROW + 1)
            if any(self.sheet.cell(row=row, column=column).value for column in (1, 2, 4, 14))
        ]
        return max(rows, default=FIRST_DATA_ROW)

    def _ensure_table(self) -> None:
        if TABLE_NAME in self.sheet.tables:
            self._resize_table(self._last_data_row())
            return
        table = Table(displayName=TABLE_NAME, ref=f"A{HEADER_ROW}:O{self._last_data_row()}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        self.sheet.add_table(table)

    def _resize_table(self, last_row: int) -> None:
        if TABLE_NAME in self.sheet.tables:
            self.sheet.tables[TABLE_NAME].ref = f"A{HEADER_ROW}:O{max(last_row, FIRST_DATA_ROW)}"

    def _add_cell_image(self, image: ExcelImage, row: int) -> None:
        self._set_cell_anchor(image, row)
        self.sheet.add_image(image)

    @staticmethod
    def _set_cell_anchor(image: ExcelImage, row: int) -> None:
        padding = 5 * 9525
        image.anchor = TwoCellAnchor(
            editAs="twoCell",
            _from=AnchorMarker(col=2, row=row - 1, colOff=padding, rowOff=padding),
            to=AnchorMarker(col=3, row=row, colOff=-padding, rowOff=-padding),
        )
